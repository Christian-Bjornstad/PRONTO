"""Complete bounded source tables; spreadsheet QC is never a review decision."""
from __future__ import annotations

import csv
from collections import Counter
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

MAX_ROWS = 10000
MAX_COLUMNS = 500
MAX_BYTES = 25 * 1024 * 1024
HEADERS = {'CNV': 'Gene_Symbol', 'RNA': 'Provisional_Event_Type', 'VARIANTS': 'Gene_symbol'}


def _scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    return str(value)


def read_source_table(path: Path, *, kind: str, demo: bool = False, sheet_name: str | None = None) -> dict:
    """Read one table, preserving ordered headers, raw values and source row IDs."""
    path = Path(path)
    if kind not in HEADERS or path.stat().st_size > MAX_BYTES:
        raise ValueError('Unsupported or oversized source table')
    payload = path.read_bytes()
    sheet = ''
    if path.suffix.lower() == '.xlsx':
        # Bound decompressed work before the workbook parser allocates cells.
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Workbook exceeds expanded size limit')
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if sheet_name is None and len(book.worksheets) != 1:
                raise ValueError('Select an explicit worksheet')
            source = book[sheet_name] if sheet_name else book.worksheets[0]
            sheet = source.title
            if source.max_column > MAX_COLUMNS or source.max_row > MAX_ROWS + 500:
                raise ValueError('Workbook exceeds row or column limit')
            source_rows = [([_scalar(cell.value) for cell in row],
                            any(cell.fill is not None and cell.fill.patternType == 'solid' and cell.fill.fgColor.type == 'rgb'
                                and str(cell.fill.fgColor.rgb)[-6:].upper() == 'FFFF00' for cell in row))
                           for row in source.iter_rows()]
        finally:
            book.close()
    elif path.suffix.lower() in {'.csv', '.tsv', '.txt'}:
        delimiter = ';' if path.suffix.lower() == '.csv' else '\t'
        with path.open(encoding='utf-8-sig', newline='') as stream:
            source_rows = []
            for row in csv.reader(stream, delimiter=delimiter):
                if len(source_rows) >= MAX_ROWS + 500 or len(row) > MAX_COLUMNS:
                    raise ValueError('Source exceeds row or column limit')
                source_rows.append((row, False))
    else:
        raise ValueError('Unsupported source table format')
    start = next((index for index, (values, _) in enumerate(source_rows)
                  if HEADERS[kind] in values), None)
    if start is None:
        raise ValueError('Expected source table header was not found')
    headers = [str(value) if value is not None else '' for value in source_rows[start][0]]
    table_id = 'source-' + sha256((kind + path.name + sheet).encode()).hexdigest()[:24]
    occurrences = Counter()
    rows = []
    for index, (values, approved) in enumerate(source_rows[start + 1:], start=start + 2):
        if not any(value not in (None, '') for value in values):
            continue
        if len(values) != len(headers):
            raise ValueError('Source row does not match table header width')
        if len(rows) >= MAX_ROWS:
            raise ValueError('Source table exceeds row limit')
        fields = dict(zip(headers, values))
        keys = {'CNV': ('Chromosome','Gene_Start','Gene_Symbol'),
                'RNA': ('Sample_ID','Provisional_Event_ID','Gene_A','Gene_B','Provisional_Event_Site_A','Provisional_Event_Site_B'),
                'VARIANTS': ('Gene_symbol','Genomic_location','DNA_change')}[kind]
        identity = [fields.get(key) for key in keys]
        digest = sha256(json.dumps([table_id, identity], sort_keys=True).encode()).hexdigest()[:24]
        finding_id = 'finding-' + digest
        occurrences[finding_id] += 1
        row = {'findingId': finding_id, 'occurrenceId': f'{finding_id}:{occurrences[finding_id]}',
               'record': index, 'values': values, 'sourceApproved': approved}
        if fields.get('Sample_ID'):
            row['sourceSampleId'] = str(fields['Sample_ID'])
        rows.append(row)
    return {'tableId': table_id, 'kind': kind, 'fileName': path.name, 'sheet': sheet,
            'sha256': sha256(payload).hexdigest(), 'headers': headers, 'rows': rows,
            'preamble': [' '.join(str(v) for v in values if v is not None)[:10000]
                         for values, _ in source_rows[:start]], 'demo': demo}


def enrich_demo_report(report, directory: Path):
    """Explicit mixed-sample demonstration input; caller must opt into demo mode."""
    from pronto_report.serialization import serialize_report_data
    from pronto_report.validation import validate_report_data
    document = json.loads(serialize_report_data(report))
    document['schemaVersion'] = '2.0'
    document.setdefault('sourceTables', [])
    for name, kind in [('CNV_summary_table.csv','CNV'),
                       ('fusion_and_splice_variant_candidates_withQC.xlsx','RNA')]:
        path = directory / name
        if path.is_file():
            table = read_source_table(path, kind=kind, demo=True)
            document['sourceTables'].append(table)
            document['provenance']['sourceFiles'].append({'name':name,'sha256':table['sha256']})
    for name in ['main_cancan_result.png','MET-splice-variant_domain_plot.pdf','TMPRSS2-ERG_2000_domain_plot.pdf']:
        path = directory / name
        if path.is_file():
            asset_id = 'example-' + sha256(name.encode()).hexdigest()[:24]
            document['attachments'].append({'assetId':asset_id, 'name': name,
                'sha256':sha256(path.read_bytes()).hexdigest(),
                'mediaType':'application/pdf' if path.suffix == '.pdf' else 'image/png',
                'description':'Mixed-sample demo source plot; not automatically matched to a table finding.'})
    for item in document['biomarkers']:
        if item['metricId'] in {'tmb','msi','hrd'}:
            item['displayRange'] = '[0-100+]' if item['metricId'] == 'tmb' else '[0-100]'
    if not any(item['metricId'] == 'hrd' for item in document['biomarkers']):
        document['biomarkers'].append({'metricId':'hrd','label':'HRD','value':42,
            'unit':'score','displayRange':'[0-100]','demo':True,
            'source':{'fileName':'synthetic-demo','field':'hrd','rawValue':42}})
    document['diagnostics'].append({'code':'MIXED_SAMPLE_DEMO','severity':'WARNING',
        'message':'Supplementary tables and plots are examples from multiple samples; not patient-matched.',
        'path':'sourceTables'})
    return validate_report_data(document)
