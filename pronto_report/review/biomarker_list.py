"""Read molecular marker names from the supplied InPreD workbook."""
from hashlib import sha256
from pathlib import Path
import re
from zipfile import ZipFile


def read_biomarker_list(path):
    path = Path(path)
    if path.suffix.lower() != '.xlsx' or path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError('Velg en biomarkørliste i XLSX-format under 10 MiB.')
    with ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
            raise ValueError('Biomarkørlisten er for stor.')
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    genes, count = set(), 0
    try:
        for sheet in book:
            if sheet.max_row > 10000 or sheet.max_column > 100:
                raise ValueError('Biomarkørlisten har for mange rader eller kolonner.')
            rows = iter(sheet.values)
            headers = next(rows, ())
            if 'Molekylær markør' not in headers or 'Biomarkør_spesifikasjon' not in headers:
                continue
            column = headers.index('Molekylær markør')
            spec_column = headers.index('Biomarkør_spesifikasjon')
            for row in rows:
                marker = str(row[column] or '').strip()
                if not marker:
                    continue
                count += 1
                if re.fullmatch(r'[A-Z][A-Z0-9-]{1,24}', marker):
                    genes.add(marker)
                elif marker == 'CDKN2A/B':
                    genes.update(('CDKN2A', 'CDKN2B'))
                elif marker in {'4q12 amplicon', '9p24.1 amplicon'}:
                    # Explicit gene symbols, excluding the workbook's PDL1/2 aliases.
                    symbols = re.findall(r'\b[A-Z][A-Z0-9]{1,24}\b', str(row[spec_column] or ''))
                    genes.update(symbol for symbol in symbols if symbol not in {'PDL1', 'PDL2'})
    finally:
        book.close()
    if not count or not genes:
        raise ValueError('Fant ikke kolonnene Molekylær markør og Biomarkør_spesifikasjon.')
    return {'genes': sorted(genes), 'rows': count, 'name': path.name,
            'sha256': sha256(path.read_bytes()).hexdigest()}
