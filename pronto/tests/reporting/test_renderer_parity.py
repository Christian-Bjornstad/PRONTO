"""Compare contract-based HTML facts with the approved legacy PPTX fixture."""

from pathlib import Path

from pptx import Presentation

from pronto.tests.reporting.test_pronto_output_adapter import build_report
from pronto_report.renderers.html import render_html


_FIXTURE_ROOT = (
    Path(__file__).parents[3]
    / "test_data"
    / "ous"
    / "251114_A02134_0115_BHCJCKDRX7_TSO_500_LocalApp_postprocessing_results"
)


def test_approved_pptx_and_html_preserve_shared_biomarker_and_case_facts():
    report = build_report()
    html = render_html(report, inline_assets=True)
    presentation = Presentation(
        _FIXTURE_ROOT / "ouput" / "IPD2225-D01-P01-A08_MTB_report.pptx"
    )
    pptx_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )

    assert report.sample["patientPseudonym"] in pptx_text
    assert report.sample["patientPseudonym"] in html
    for biomarker in report.biomarkers:
        assert biomarker["source"]["rawValue"] in pptx_text
    assert "14,9 mut/Mb" in html
    assert "4,13 %" in html
    assert len(presentation.slides) == 12
    assert len(report.variants) == 30
