from importlib import import_module, util

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill


def reader():
    assert util.find_spec('pronto_report.adapters.source_tables'), 'complete table importer missing'
    return import_module('pronto_report.adapters.source_tables').read_source_table


def test_cnv_preserves_preamble_empty_columns_and_all_values(tmp_path):
    file = tmp_path / 'cnv.csv'
    file.write_text('# source;;;;\nChromosome;Gene_Start;Gene_Symbol;Tumor_CN;Extra\nchr7;100;DEMO;3.14;\n')
    table = reader()(file, kind='CNV')
    assert table['headers'] == ['Chromosome','Gene_Start','Gene_Symbol','Tumor_CN','Extra']
    assert table['rows'][0]['values'][-1] == ''
    assert table['rows'][0]['values'][3] == '3.14'
    assert table['preamble'][0].startswith('# source')


def test_rna_retains_formula_text_source_qc_and_duplicate_headers(tmp_path):
    file = tmp_path/'rna.xlsx'
    book = Workbook(); sheet = book.active
    sheet.append(['# Legend'])
    sheet.append(['Provisional_Event_Type','Sample_ID','Gene_A','Gene_B','QC_Verdict','Extra','Extra'])
    sheet.append(['Fusion','DEMO-RNA','A','B','OK','=1+2',None])
    sheet['A3'].fill = PatternFill('solid',fgColor='FFFF00')
    book.save(file)
    table = reader()(file, kind='RNA')
    assert table['headers'][-2:] == ['Extra','Extra']
    assert table['rows'][0]['values'][-2:] == ['=1+2',None]
    assert table['rows'][0]['sourceApproved'] is True
    assert table['rows'][0]['sourceSampleId'] == 'DEMO-RNA'
    assert table['rows'][0]['record'] == 3
    assert table['rows'][0]['values'][4] == 'OK'


def test_unsupported_source_is_rejected_without_guessing(tmp_path):
    file = tmp_path/'unknown.csv'; file.write_text('a;b\n1;2\n')
    with pytest.raises(ValueError, match='header'):
        reader()(file,kind='CNV')
